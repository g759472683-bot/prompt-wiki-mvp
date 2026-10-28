from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from app.services.parsing import parse_document


def _make_pdf(text: str) -> bytes:
    """Build a tiny one-page PDF with a built-in Helvetica font."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_parse_pdf(tmp_path: Path):
    path = tmp_path / "sample.pdf"
    path.write_bytes(_make_pdf("Hello backend PDF"))
    result = parse_document(path, "sample.pdf")
    assert result.parser == "pdfplumber"
    assert result.pages == 1
    assert result.blocks
    assert "Hello backend PDF" in " ".join(block.text for block in result.blocks)


def test_parse_docx(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("第一段：账务平台年度可用性 99.99%。")
    doc.add_paragraph("第二段：对账延迟 P95 从 45 分钟降至 8 分钟。")
    doc.save(path)

    result = parse_document(path, "sample.docx")
    assert result.parser == "python-docx"
    assert len(result.blocks) == 2
    assert "99.99%" in result.blocks[0].text


def test_parse_txt_keeps_offsets(tmp_path: Path):
    path = tmp_path / "sample.txt"
    path.write_text("第一段\n\n第二段\n\n第三段", encoding="utf-8")
    result = parse_document(path, "sample.txt")
    assert len(result.blocks) == 3
    assert result.blocks[0].start == 0
    assert result.blocks[1].start > result.blocks[0].end
